"""Build CLAIM_EVIDENCE_MATRIX.xlsx: each substantive claim in Chapters 4-6 -> evidence -> status."""

import os

import openpyxl
from openpyxl.styles import Alignment, Font, PatternFill

R = "results/pilot/"
CLAIMS = [
    # id, chapter/section, claim, evidence, status, caveat
    ("C4-1", "4.10.2", "REM runs inside AgentDojo's execution loop without modifying the agent (AS-1)", "rem/integration/agentdojo_replay.py; tests/test_replay.py (3 passed)", "SUPPORTED (executed)", "Demonstrated with replayed, not live, agents"),
    ("C4-2", "4.12", "192 tests pass", "pytest run 27 Oct 2026, Python 3.11.17", "SUPPORTED (executed)", "189 pass and the replay module is skipped without AgentDojo"),
    ("C4-3", "4.12", "Pilot reproduces exactly from commit aa66f22", "experiments/pilot/run_all.sh; two full runs, identical outputs", "SUPPORTED (executed)", ""),
    ("C4-4", "4.9", "Held-out predictions are independent of held-out labels", "tests/test_crossfit.py::test_crossfit_no_leakage_and_outputs", "SUPPORTED (executed)", "Unit test on synthetic data"),
    ("C4-5", "4.7.7", "Attribution cannot reach the verdict", "tests/test_contract_compliance.py::test_decision_engine_signature_excludes_attribution", "SUPPORTED (executed)", ""),
    ("C4-6", "4.11", "A2 issues no verdict", "tests/test_a2.py::test_a2_is_observe_only_and_issues_no_verdict", "SUPPORTED (executed)", ""),
    ("C4-7", "4.10.3", "Replay reproduces 3,313/3,334 recorded outcomes", R + "e4_rmod_and_replay_validity.json: replay_agreement_total", "SUPPORTED (executed)", "Suite v1 used for all replays"),
    ("C4-8", "4.5", "x_ctx1 classifier not implemented", "PROJECT_AUDIT.md; network refusal of model host", "BLOCKED", "SD-DET"),
    ("C5-1", "5.2.1", "3,334 episodes, 6,762 steps, 921 positive", R + "dataset_summary.json", "SUPPORTED (executed)", "Not i.i.d."),
    ("C5-2", "5.3", "All evidence groups pass E0 (ctx 0.582, beh 0.802, act 0.929)", R + "e0_pilot.json", "SUPPORTED (exploratory)", "Context marginal"),
    ("C5-3", "5.4", "Held-out ROC-AUC 0.938 [0.908, 0.961]", R + "e1_crossfit.json: label_nll_primary.ci_roc_auc", "SUPPORTED (exploratory)", "9 groups; one benchmark"),
    ("C5-4", "5.4", "Calibration gives no measurable ECE gain (0.044 vs 0.042)", R + "e1_crossfit.json: pooled_calibrated vs pooled_uncalibrated", "SUPPORTED (exploratory; negative)", ""),
    ("C5-5", "5.4", "Calibration degrades on unseen credential-change attacks (fold 7 ECE 0.189)", R + "e1_crossfit.json: per_fold.injection_task_7", "SUPPORTED (exploratory)", "Single fold"),
    ("C5-6", "5.4", "Cross-template ROC-AUC 0.92-0.96", R + "e4_cross_template_and_defenses.json: cross_template.*.steps", "SUPPORTED (exploratory)", "gpt-4o only"),
    ("C5-7", "5.4", "Leave-one-backbone-out ROC-AUC 0.949 pooled; 0.736-1.000 per backbone", R + "e1_crossfit.json: grouping_pipeline", "SUPPORTED (exploratory)", "Few positives for some backbones"),
    ("C5-8", "5.5", "REM step FPR 0.227 vs B1 0.056 (observe-only)", R + "e2_e5_e6_observe_only.json: verdicts.*.ci_step_fpr", "SUPPORTED (exploratory)", "Mostly Escalate"),
    ("C5-9", "5.5", "REM releases 33/180 benign state-changing steps; B1 132/180", R + "benign_release_breakdown.json", "SUPPORTED (exploratory)", ""),
    ("C5-10", "5.6", "REM ASR 0.004 vs B0 0.226 vs B1 0.046-0.055 (replay)", R + "e4_e5_enforcement_replay.json: configs", "OUT OF SCOPE (SC-7) / exploratory", "NO-ADAPT"),
    ("C5-11", "5.6", "REM benign utility 0.314 (no reviewer) / 0.497 (idealized)", R + "e4_e5_enforcement_replay.json", "OUT OF SCOPE (SC-7) / exploratory", "Oracle reviewer is an upper bound"),
    ("C5-12", "5.6", "r_mod = 0.000 (tier 3), 0.010 (tier 4)", R + "e4_rmod_and_replay_validity.json", "OUT OF SCOPE (SC-7)", "Artefact of NO-ADAPT"),
    ("C5-13", "5.6", "Latency median 0.17-0.20 ms, p95 0.43-0.54 ms; memory median 2.9 KB", R + "e4_e5_enforcement_replay.json; latency_memory_profile.json", "SUPPORTED (executed)", "Container hardware; excludes LLM and tools"),
    ("C5-14", "5.7", "Removing behavioral evidence raises benign utility 0.314 -> 0.419 at equal ASR", R + "e4_e5_enforcement_replay.json: abl_no_beh_*", "OUT OF SCOPE (SC-7) / exploratory; negative", ""),
    ("C5-15", "5.7", "Removing action evidence lowers ROC-AUC to 0.799", R + "e1_crossfit.json: ablation_refits.no_act", "SUPPORTED (exploratory)", ""),
    ("C5-16", "5.8", "A2 held-out benign false alarm 0.106 vs target 0.05", R + "e3_a2_cusum_observe_only.json", "SUPPORTED (exploratory; negative)", ""),
    ("C5-17", "5.9", "user_task_15 collision in suite v1/v1.1; fixed upstream from v1.1.1", "agentdojo 0.1.35 package ground truths (v1..v1.2.2); session check 27 Oct 2026", "SUPPORTED (checked)", "Not a new finding"),
    ("C5-18", "5.9", "Post-hoc: excluding user_task_15 gives REM ASR 0.000", R + "sensitivity_post_hoc_exclude_user_task_15.json", "SUPPORTED (post-hoc, disclosed)", "Not pre-registered"),
    ("C5-19", "5.9", "Label audit: 29/30 consistent", R + "label_audit_sample.md", "PARTIAL", "Inspected by AI assistant, not a human"),
    ("C5-20", "5.10", "REM below B1 in all 36 grid cells, matched and misspecified", R + "e2_e5_e6_observe_only.json: e6_grid", "SUPPORTED (exploratory)", "Favoured by construction"),
    ("C6-1", "6.1.1", "Runtime cost is not an obstacle to step-level mediation", "C5-13", "SUPPORTED (exploratory)", "Excludes reviewer latency"),
    ("C6-2", "6.1.2", "REM as configured is not deployable without a reviewer model", "C5-11", "SUPPORTED (exploratory) / SC-7", ""),
    ("C6-3", "6.4", "Contributions are compositional / evaluative, not methodological", "RESEARCH_GAP_AND_NOVELTY_AUDIT.md; Chapter 3 Table 3.16", "SUPPORTED (argued)", "Bounded by the Chapter 2 search"),
    ("C6-4", "6.3", "REM differs in design from NEXUS, AgentTrust, SafeAgent, etc.", "Chapter 2 review; STAGE2_5_REM_VS_PRIOR_ART_MATRIX.md", "SUPPORTED (design-level only)", "NEXUS reading OI-01 unresolved; no shared evaluation"),
]


def main():
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = "Claims"
    ws.append(["ID", "Section", "Claim", "Evidence (file / test / key)", "Status", "Caveat"])
    for c in ws[1]:
        c.font = Font(bold=True)
    for row in CLAIMS:
        ws.append(list(row))
    colors = {"SUPPORTED": "E3F1E7", "OUT OF SCOPE": "FBE3E1", "BLOCKED": "FBE3E1", "PARTIAL": "FFF4D6"}
    for row in ws.iter_rows(min_row=2):
        for k, col in colors.items():
            if (row[4].value or "").startswith(k):
                row[4].fill = PatternFill("solid", fgColor=col)
        for c in row:
            c.alignment = Alignment(wrap_text=True, vertical="top")
    for col, w in zip("ABCDEF", (8, 9, 60, 55, 30, 36)):
        ws.column_dimensions[col].width = w
    root = os.path.abspath(os.path.join(os.path.dirname(__file__), ".."))
    wb.save(os.path.join(root, "CLAIM_EVIDENCE_MATRIX.xlsx"))
    print(len(CLAIMS), "claims")


if __name__ == "__main__":
    main()
