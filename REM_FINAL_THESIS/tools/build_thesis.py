"""Build the final thesis: per-chapter DOCX, full DOCX, and a PDF rendered by LibreOffice.
usage: python3 REM_FINAL_THESIS/tools/build_thesis.py   (from the repository root)
Requires pandoc >= 3.5 (pypandoc_binary), python-docx and LibreOffice (soffice).
Formatting is provisional: the university template was not available."""
import os, re, subprocess, sys, time, zipfile, shutil
import pypandoc
from docx import Document
from docx.shared import Pt, Cm, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

ROOT = os.path.abspath(os.path.join(os.path.dirname(__file__), '..', '..'))
CH = os.path.join(ROOT, 'REM_FINAL_THESIS', 'chapters')
OUT = os.path.join(ROOT, 'REM_FINAL_THESIS')
BUILD = os.path.join(OUT, 'build'); os.makedirs(BUILD, exist_ok=True)
P = pypandoc.get_pandoc_path()
FMT = 'markdown+lists_without_preceding_blankline'
CHAPTERS = ['CHAPTER1_INTRODUCTION', 'CHAPTER2_LITERATURE_REVIEW', 'CHAPTER3_METHODOLOGY',
            'CHAPTER4_IMPLEMENTATION', 'CHAPTER5_RESULTS', 'CHAPTER6_DISCUSSION_CONCLUSION']

# ------------------------------------------------------------------ reference document (styles)
ref0 = os.path.join(BUILD, 'pandoc_default_reference.docx')
with open(ref0, 'wb') as f:
    f.write(subprocess.run([P, '--print-default-data-file', 'reference.docx'], capture_output=True, check=True).stdout)
doc = Document(ref0)
FONT = 'Times New Roman'
def font(style, size, bold=None, italic=None):
    rpr = style.element.get_or_add_rPr()
    rf = rpr.find(qn('w:rFonts'))
    if rf is None:
        rf = OxmlElement('w:rFonts'); rpr.insert(0, rf)
    for a in list(rf.attrib): del rf.attrib[a]
    for a in ('w:ascii', 'w:hAnsi', 'w:cs', 'w:eastAsia'): rf.set(qn(a), FONT)
    style.font.size = Pt(size); style.font.color.rgb = RGBColor(0, 0, 0)
    if bold is not None: style.font.bold = bold
    if italic is not None: style.font.italic = italic
st = doc.styles
for n in ('Normal', 'Body Text', 'First Paragraph', 'Compact'):
    try:
        font(st[n], 12); pf = st[n].paragraph_format; pf.line_spacing = 1.5 if n != 'Compact' else 1.15
        pf.space_after = Pt(6)
        if n in ('Body Text', 'First Paragraph'): pf.alignment = WD_ALIGN_PARAGRAPH.JUSTIFY
    except KeyError:
        pass
for n, sz in (('Heading 1', 16), ('Heading 2', 14), ('Heading 3', 12), ('Heading 4', 12)):
    font(st[n], sz, bold=True, italic=(n == 'Heading 4')); st[n].paragraph_format.keep_with_next = True
    st[n].paragraph_format.space_before = Pt(18 if n == 'Heading 1' else 12); st[n].paragraph_format.space_after = Pt(8)
st['Heading 1'].paragraph_format.page_break_before = True
for n, sz in (('Title', 20), ('Subtitle', 12), ('Author', 12), ('Date', 12)):
    try:
        font(st[n], sz, bold=(n == 'Title'), italic=(n == 'Subtitle'))
        st[n].paragraph_format.alignment = WD_ALIGN_PARAGRAPH.CENTER; st[n].paragraph_format.space_after = Pt(12)
    except KeyError:
        pass
for n in ('Table Caption', 'Image Caption', 'Caption'):
    try:
        font(st[n], 11, bold=(n == 'Table Caption'), italic=(n != 'Table Caption'))
    except KeyError:
        pass
for n in ('Source Code', 'Verbatim Char'):
    try:
        st[n].font.size = Pt(9.5)
    except KeyError:
        pass
# compact table text
try:
    tbl = st['Table']
    ppr = tbl.element.get_or_add_pPr()
except KeyError:
    pass
sec = doc.sections[0]
sec.page_width, sec.page_height = Cm(21.0), Cm(29.7)
sec.left_margin = Cm(3.0); sec.right_margin = Cm(2.5); sec.top_margin = Cm(2.5); sec.bottom_margin = Cm(2.5)
fp = sec.footer.paragraphs[0]; fp.text = ''; fp.alignment = WD_ALIGN_PARAGRAPH.CENTER
for kind, txt in (('begin', None), ('instr', ' PAGE '), ('separate', None), ('text', '1'), ('end', None)):
    r = fp.add_run()
    if kind == 'instr':
        e = OxmlElement('w:instrText'); e.set(qn('xml:space'), 'preserve'); e.text = txt; r._r.append(e)
    elif kind == 'text':
        r.text = txt
    else:
        e = OxmlElement('w:fldChar'); e.set(qn('w:fldCharType'), kind); r._r.append(e)
doc.core_properties.title = 'Runtime Evaluation and Mitigation (REM): A Security Framework for Financial AI Agents'
doc.core_properties.author = ''
REF = os.path.join(BUILD, 'reference.docx'); doc.save(REF)

# ------------------------------------------------------------------ source preparation
def prepare(md):
    """Bold table titles directly above a pipe table become pandoc table captions (for the list of tables)."""
    md = re.sub(r'^# (CHAPTER [A-Z]+)\n\n?# (.+)$', lambda m: f'# {m.group(1)}: {m.group(2)}', md, count=1, flags=re.M)
    L = md.split('\n'); out = []; i = 0
    while i < len(L):
        m = re.match(r'^\*\*(Table \d+[A-Z]?\.\d+\.?.*?)\*\*(.*)$', L[i])
        if m:
            j = i + 1
            while j < len(L) and not L[j].strip(): j += 1
            if j < len(L) and L[j].startswith('|'):
                cap = (m.group(1) + m.group(2)).strip()
                out.append(f'Table: {cap}'); out.append(''); i += 1; continue
        out.append(L[i]); i += 1
    return '\n'.join(out)

tmp = os.path.join(BUILD, 'src'); os.makedirs(tmp, exist_ok=True)
srcs = []
for name in ['FRONT_MATTER'] + [c + '_FINAL' for c in CHAPTERS] + ['REFERENCES']:
    t = open(os.path.join(CH, name + '.md'), encoding='utf-8').read()
    if name == 'REFERENCES':
        t = t.replace('# REFERENCES', '# References', 1)
    p = os.path.join(tmp, name + '.md'); open(p, 'w', encoding='utf-8').write(prepare(t)); srcs.append(p)

def pandoc(inputs, out, extra=()):
    subprocess.run([P, '-f', FMT, *inputs, '--resource-path', ROOT, '--reference-doc', REF, '-o', out, *extra], check=True, cwd=ROOT)

for c, s in zip(CHAPTERS, srcs[1:7]):
    pandoc([s, srcs[-1]], os.path.join(OUT, c + '_FINAL.docx'))
FULL = os.path.join(OUT, 'REM_FINAL_THESIS.docx')
pandoc(srcs, FULL, ['--toc', '--toc-depth=2', '--lof', '--lot'])

# Word: refresh fields (TOC, lists, page numbers) on open
z = zipfile.ZipFile(FULL); F = {n: z.read(n) for n in z.namelist()}; z.close()
s = F['word/settings.xml'].decode()
if 'w:updateFields' not in s:
    s = s.replace('</w:settings>', '<w:updateFields w:val="true"/></w:settings>')
F['word/settings.xml'] = s.encode()
with zipfile.ZipFile(FULL, 'w', zipfile.ZIP_DEFLATED) as zo:
    for n in sorted(F, key=lambda n: (n != '[Content_Types].xml', n)): zo.writestr(n, F[n])

# ------------------------------------------------------------------ PDF via LibreOffice, lists of figures/tables filled
import html, json
import uno
from com.sun.star.beans import PropertyValue
def pv(n, v):
    p = PropertyValue(); p.Name = n; p.Value = v; return p
PDF = os.path.join(OUT, 'REM_FINAL_THESIS.pdf')

def export_pdf():
    prof = os.path.join(BUILD, 'lo_profile')
    proc = subprocess.Popen(['soffice', '-env:UserInstallation=file://' + prof, '--headless', '--invisible', '--norestore',
                             '--nologo', '--accept=socket,host=127.0.0.1,port=2019;urp;'],
                            stdout=subprocess.DEVNULL, stderr=subprocess.DEVNULL, env=dict(os.environ, SAL_USE_VCLPLUGIN='svp'))
    local = uno.getComponentContext()
    res = local.ServiceManager.createInstanceWithContext('com.sun.star.bridge.UnoUrlResolver', local)
    ctx = None
    for _ in range(90):
        try:
            ctx = res.resolve('uno:socket,host=127.0.0.1,port=2019;urp;StarOffice.ComponentContext'); break
        except Exception:
            time.sleep(1)
    assert ctx is not None, 'LibreOffice did not start'
    desk = ctx.ServiceManager.createInstanceWithContext('com.sun.star.frame.Desktop', ctx)
    d = desk.loadComponentFromURL('file://' + FULL, '_blank', 0, (pv('Hidden', True),))
    idx = d.getDocumentIndexes()
    if idx.getCount():
        idx.getByIndex(0).update()          # table of contents (heading-based)
    d.refresh()
    d.storeToURL('file://' + PDF, (pv('FilterName', 'writer_pdf_Export'),))
    d.close(True)
    try:
        desk.terminate()
    except Exception:
        pass
    proc.wait(timeout=60)

def read_docx():
    z = zipfile.ZipFile(FULL); F = {n: z.read(n) for n in z.namelist()}; z.close(); return F
def write_docx(F):
    with zipfile.ZipFile(FULL, 'w', zipfile.ZIP_DEFLATED) as zo:
        for n in sorted(F, key=lambda n: (n != '[Content_Types].xml', n)): zo.writestr(n, F[n])

F = read_docx(); X = F['word/document.xml'].decode()
def captions(style):
    out = []
    for m in re.finditer(r'<w:p>(?:(?!</w:p>).)*?<w:pStyle w:val="%s" />.*?</w:p>' % style, X, flags=re.S):
        out.append(html.unescape(''.join(re.findall(r'<w:t[^>]*>([^<]*)</w:t>', m.group(0)))).strip())
    return out
CAPS = {'Image Caption': captions('ImageCaption'), 'Table Caption': captions('TableCaption')}

def pages_of(caps):
    txt = subprocess.run(['pdftotext', '-layout', PDF, '-'], capture_output=True, text=True).stdout.split('\f')
    norm = lambda t: re.sub(r'\s+', ' ', t)
    pages = [norm(p) for p in txt]
    start = next(i for i, p in enumerate(pages) if 'Status of This Document' in p and 'Table of Contents' not in p)
    res = []
    for c in caps:
        key = norm(c)[:55]
        pg = next((i + 1 for i in range(start, len(pages)) if key in pages[i]), None)
        res.append(pg)
    return res

def fill(pages):
    global X
    TW = int((21.0 - 3.0 - 2.5) / 2.54 * 1440)
    for field, style in (('Image Caption', 'Image Caption'), ('Table Caption', 'Table Caption')):
        paras = []
        for k, (c, pg) in enumerate(zip(CAPS[style], pages[style])):
            start = ('<w:r><w:fldChar w:fldCharType="begin" w:dirty="true" /></w:r><w:r><w:instrText xml:space="preserve">TOC \\h \\z \\t &quot;%s&quot; \\c</w:instrText></w:r><w:r><w:fldChar w:fldCharType="separate" /></w:r>' % field) if k == 0 else ''
            end = '<w:r><w:fldChar w:fldCharType="end" /></w:r>' if k == len(CAPS[style]) - 1 else ''
            paras.append('<w:p><w:pPr><w:tabs><w:tab w:val="right" w:leader="dot" w:pos="%d" /></w:tabs><w:spacing w:after="60" w:line="240" w:lineRule="auto" /><w:ind w:left="0" w:hanging="0" /></w:pPr>%s<w:r><w:rPr><w:sz w:val="22" /></w:rPr><w:t xml:space="preserve">%s</w:t></w:r><w:r><w:tab /></w:r><w:r><w:rPr><w:sz w:val="22" /></w:rPr><w:t>%s</w:t></w:r>%s</w:p>'
                         % (TW, start, html.escape(c, quote=False), pg if pg else '', end))
        pat = re.compile(r'<w:p>(?:(?!<w:p>).)*?TOC \\h \\z \\t &quot;%s&quot; \\c.*?</w:p>(?:<w:p><w:pPr><w:tabs>.*?</w:p>)*' % field.replace(' ', ' '), flags=re.S)
        m = pat.search(X)
        assert m, field
        # replace the field paragraph and any previously filled entries up to the end of the sdtContent
        end_idx = X.index('</w:sdtContent>', m.start())
        X = X[:m.start()] + ''.join(paras) + X[end_idx:]
    F['word/document.xml'] = X.encode(); write_docx(F)

pages = {k: [0] * len(v) for k, v in CAPS.items()}
fill(pages); export_pdf()
for _ in range(2):
    new = {k: pages_of(v) for k, v in CAPS.items()}
    if new == pages: break
    pages = new; fill(pages); export_pdf()
missing = {k: sum(1 for p in v if not p) for k, v in pages.items()}
json.dump({k: list(zip(CAPS[k], v)) for k, v in pages.items()}, open(os.path.join(BUILD, 'lists_pages.json'), 'w'), indent=1)
print('built:', FULL, os.path.getsize(FULL), 'bytes; PDF', os.path.getsize(PDF), 'bytes; figures', len(CAPS['Image Caption']),
      'tables', len(CAPS['Table Caption']), 'captions without page', missing)
